"""Audit-fix lane L4 (ingestion): A1-01 A1-06 A1-03 A1-07 A1-08 A1-10."""

import zlib
from unittest.mock import AsyncMock, MagicMock

import numpy as np
import pytest

from app.indexing import service as idx


class _Emb:
    async def embed_texts(self, texts, batch_size=None, progress_callback=None):
        if progress_callback:
            progress_callback(1, 1)
        return np.array([[float(i + 1)] for i, _ in enumerate(texts)], dtype=np.float32)


class _Lance:
    def __init__(self):
        self.summaries = []
        self.deleted = []

    async def delete_documents(self, ids):
        self.deleted.append(ids)

    async def add_documents(self, ids, embs, metas):
        pass

    async def add_summaries_batch(self, items):
        self.summaries.extend(items)

    async def clear_query_cache(self):
        pass

    async def create_hnsw_index(self, table):
        pass

    def get_max_id(self, table_name="pma_chunks"):
        return 0


async def _service(tmp_path):
    from app.storage.db import DatabaseManager

    mgr = DatabaseManager(str(tmp_path / "l4.db"))
    await mgr.init_db(schema_path="app/storage/schema.sql")
    return mgr, idx.IndexingService(mgr, _Emb(), _Lance())


async def _chunk_texts(mgr):
    rows = await mgr.execute_query("SELECT text_preview FROM chunks")
    return [zlib.decompress(r[0]).decode("utf-8") for r in rows]


# --- A1-01 ---------------------------------------------------------------


@pytest.mark.asyncio
async def test_a1_01_streamed_extractor_records_are_newline_separated(tmp_path):
    import docx
    import openpyxl

    corpus = tmp_path / "docs"
    corpus.mkdir()
    (corpus / "people.csv").write_text("name,city\nAlice,Paris\nBob,Oslo\n", encoding="utf-8")
    d = docx.Document()
    d.add_heading("Quarterly Plan", level=1)
    d.add_paragraph("We will ship the beta in March.")
    d.add_heading("Risks", level=2)
    d.add_paragraph("Hiring is slow.")
    d.save(str(corpus / "plan.docx"))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Item", "Cost"])
    ws.append(["Rent", 1200])
    ws.append(["Food", 300])
    wb.save(str(corpus / "budget.xlsx"))

    mgr, service = await _service(tmp_path)
    idx.progress.status = "idle"
    try:
        await service.index_folders([str(corpus)])
        text = "\n".join(await _chunk_texts(mgr))
    finally:
        await mgr.close()

    assert "Paris\nname: Bob" in text
    assert "March.\n## Risks\n" in text
    assert "Rent | 1200\nFood | 300" in text
    assert "Parisname" not in text


# --- A1-06 ---------------------------------------------------------------


@pytest.mark.asyncio
async def test_a1_06_pre_extract_group_is_bounded_by_bytes(tmp_path, monkeypatch):
    seen = []
    fake = MagicMock()
    fake.extract_text_files = lambda paths, cap: seen.append(list(paths)) or []
    monkeypatch.setattr(idx, "rust_core", fake)
    monkeypatch.setattr(idx, "RUST_CORE_AVAILABLE", True)
    monkeypatch.setattr(idx, "_PRE_EXTRACT_FILE_MAX_BYTES", 100)
    monkeypatch.setattr(idx, "_PRE_EXTRACT_GROUP_MAX_BYTES", 250)

    files = []
    for name, size in (("a.txt", 90), ("huge.log", 5000), ("b.txt", 90), ("c.txt", 90)):
        p = tmp_path / name
        p.write_bytes(b"x" * size)
        files.append((p, "t", str(tmp_path)))

    mgr, service = await _service(tmp_path)
    try:
        await service._rust_pre_extract(files)
    finally:
        await mgr.close()

    # huge.log is over the per-file cap; c.txt would take the group past 250.
    assert [p.rsplit("\\", 1)[-1] for p in seen[0]] == ["a.txt", "b.txt"]


# --- A1-03 ---------------------------------------------------------------


@pytest.mark.asyncio
async def test_a1_03_cancelled_run_still_flushes_summary_vectors(tmp_path):
    corpus = tmp_path / "docs"
    corpus.mkdir()
    (corpus / "a.txt").write_text("hello world " * 20, encoding="utf-8")

    mgr, service = await _service(tmp_path)
    ids = await mgr.batch_insert_files(
        [
            {
                "path": str(corpus / "a.txt"),
                "size": 1,
                "modified_at": "2026-01-01T00:00:00",
                "type": ".txt",
                "folder_tag": "docs",
                "sha256": "b" * 64,
            }
        ]
    )
    await mgr.execute_write("UPDATE files SET summary = ? WHERE id = ?", ("A summary.", ids[0]))

    async def cancelled_pipeline(*_a, **_k):
        service._summary_dirty_file_ids.add(ids[0])
        idx.progress.is_cancelled = True

    service._batch_index_pipeline = cancelled_pipeline
    idx.progress.status = "idle"
    try:
        await service.index_folders([str(corpus)])
    finally:
        idx.progress.is_cancelled = False  # module-level singleton: don't leak
        await mgr.close()

    written = [s["doc_id"] for s in service.lancedb_client.summaries]
    assert f"file_{ids[0]}" in written


# --- A1-07 ---------------------------------------------------------------


@pytest.mark.asyncio
async def test_a1_07_watcher_skips_folders_with_no_indexed_files(tmp_path):
    from app.indexing.watcher import FolderWatcher

    kept, removed = tmp_path / "kept", tmp_path / "removed"
    # Sibling whose name starts with `removed`: its files must not keep
    # `removed` alive (bare-prefix match).
    sibling = tmp_path / "removed2"
    for d in (kept, removed, sibling):
        d.mkdir()
    mgr, _ = await _service(tmp_path)
    try:
        for folder in (kept, removed, sibling):
            await mgr.upsert_folder_profile(
                {
                    "folder_path": str(folder),
                    "folder_tag": folder.name,
                    "profile_text": "p",
                    "project_type": "docs",
                    "file_count": 1,
                    "total_size_bytes": 1,
                    "top_extensions": "[]",
                    "key_files": "[]",
                }
            )
        # Only `kept` still has files: `removed` models a folder whose index
        # was removed (files deleted, folder_profiles row left behind).
        await mgr.batch_insert_files(
            [
                {
                    "path": str(kept / "a.txt"),
                    "size": 1,
                    "modified_at": "2026-01-01T00:00:00",
                    "type": ".txt",
                    "folder_tag": "kept",
                    "sha256": "b" * 64,
                },
                {
                    "path": str(sibling / "a.txt"),
                    "size": 1,
                    "modified_at": "2026-01-01T00:00:00",
                    "type": ".txt",
                    "folder_tag": "removed2",
                    "sha256": "c" * 64,
                },
            ]
        )
        service = MagicMock(index_folders=AsyncMock())
        watcher = FolderWatcher(mgr, lambda: service)
        roots = await watcher.run_once()
    finally:
        await mgr.close()

    assert sorted(roots) == sorted([str(kept), str(sibling)])


# --- A1-08 ---------------------------------------------------------------


@pytest.mark.asyncio
async def test_a1_08_folder_profile_vectors_are_replaced_not_appended(tmp_path):
    corpus = tmp_path / "recipes"
    corpus.mkdir()
    (corpus / "a.txt").write_text("hello world " * 20, encoding="utf-8")
    mgr, service = await _service(tmp_path)

    store: dict[str, dict] = {}

    async def add(items):
        for s in items:
            # A plain append, like LanceDB's tbl.add: collect, never overwrite.
            store.setdefault(s["doc_id"], {"n": 0})["n"] += 1

    async def delete(ids):
        for i in ids:
            store.pop(i, None)

    service.lancedb_client.add_summaries_batch = add
    service.lancedb_client.delete_summaries_by_ids = delete
    all_files = [(corpus / "a.txt", "recipes", str(corpus))]
    try:
        for _ in range(3):
            await service._generate_folder_profiles(all_files, [corpus])
    finally:
        await mgr.close()

    assert {k: v["n"] for k, v in store.items()} == {"folder_profile_recipes": 1}


# --- A1-10 ---------------------------------------------------------------


@pytest.mark.asyncio
async def test_a1_10_ocr_chunks_do_not_overlap_native_offsets(tmp_path):
    from pathlib import Path

    from app.ocr.types import OcrLine, OcrPage

    path = Path(r"C:\docs\mixed.pdf")
    mgr, service = await _service(tmp_path)
    service.lancedb_client.delete_documents = AsyncMock()
    try:
        ids = await mgr.batch_insert_files(
            [
                {
                    "path": str(path.absolute()),
                    "size": 1,
                    "modified_at": "2026-01-01T00:00:00",
                    "type": ".pdf",
                    "folder_tag": "docs",
                    "sha256": "b" * 64,
                }
            ]
        )
        await mgr.insert_chunks_bulk(
            [
                {
                    "file_id": ids[0],
                    "start_offset": 0,
                    "end_offset": 1004,
                    "text_preview": "native body text",
                    "sentence_offsets": "[]",
                    "segmenter_version": "py_v1",
                    "source": None,
                }
            ]
        )
        page = OcrPage(
            page_num=0,
            lines=(OcrLine(text="Scanned appendix signed by Okafor. " * 4, conf=0.95, low=False),),
            mean_conf=0.9,
        )
        idx.progress.status = "idle"
        await service.index_ocr_pages(path, [page])
        rows = await mgr.execute_query(
            "SELECT start_offset FROM chunks WHERE file_id = ? AND source = 'ocr'", (ids[0],)
        )
    finally:
        await mgr.close()

    assert rows and min(r[0] for r in rows) >= 1004


# --- A1-12 ---------------------------------------------------------------


def test_a1_12_xlsx_zip_bomb_is_rejected_before_openpyxl(tmp_path, monkeypatch):
    import zipfile

    import openpyxl

    from app.indexing.extractors.xlsx_extractor import XlsxExtractor

    bomb = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/sharedStrings.xml", b"\0" * (3 * 1024 * 1024))  # ratio >> 200

    opened = []
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **k: opened.append(a))
    assert list(XlsxExtractor().extract_stream(bomb, 50_000_000)) == []
    assert opened == []


# --- A1-13 ---------------------------------------------------------------


def test_a1_13_json_is_not_cut_at_200k_chars(tmp_path):
    import json

    from app.indexing.extractors.json_extractor import JsonExtractor

    path = tmp_path / "customers.json"
    path.write_text(
        json.dumps([{"id": i, "name": f"cust{i}"} for i in range(9000)] + ["LAST-RECORD-MARKER"]),
        encoding="utf-8",
    )
    assert path.stat().st_size < 500_000
    out = "".join(JsonExtractor().extract_stream(path, 50_000_000))
    assert len(out) > 200_000
    assert "LAST-RECORD-MARKER" in out
    # The cap is max_file_size, not a hidden constant.
    assert len("".join(JsonExtractor().extract_stream(path, 4000))) <= 4000


def test_a1_13_plain_text_fallback_honours_max_file_size(tmp_path):
    (tmp_path / "dump.sql").write_text("INSERT INTO t VALUES (1);\n" * 2000, encoding="utf-8")
    svc = idx.IndexingService(MagicMock(), _Emb(), _Lance())
    svc.max_file_size = 4000
    assert sum(len(c) for c in svc._extract_plain_text_stream(tmp_path / "dump.sql")) == 4000


# --- A1-14 ---------------------------------------------------------------


def test_a1_14_python_chunks_keep_decorators_and_leading_comments():
    from app.indexing.code_chunker import CodeChunker

    src = (
        "import os\n"
        "\n"
        "# INC-4471: webhook must stay idempotent\n"
        '@app.post("/billing/webhook")\n'
        '@requires_role("finance-admin")\n'
        "def hook():\n"
        "    return 1\n"
        "\n"
        "@dataclass(frozen=True)\n"
        "class Point:\n"
        "    x: int = 0\n"
    )
    chunks = CodeChunker().chunk_code(src, "m.py")
    joined = "\n".join(c["text_preview"] for c in chunks)
    for needle in ('@app.post("/billing/webhook")', "@requires_role", "INC-4471", "@dataclass"):
        assert needle in joined, needle
    # Offsets still address the source: each span's text matches its chunk body.
    for c in chunks:
        assert src[c["start_offset"] : c["end_offset"]].strip() in c["text_preview"]


# --- A1-15 ---------------------------------------------------------------


@pytest.mark.asyncio
@pytest.mark.parametrize("cancel_in", ["_scan_all_folders", "_detect_changes"])
async def test_a1_15_cancel_during_scan_or_detection_is_honoured(tmp_path, cancel_in):
    corpus = tmp_path / "docs"
    corpus.mkdir()
    (corpus / "a.txt").write_text("hello world " * 20, encoding="utf-8")

    mgr, service = await _service(tmp_path)
    real = getattr(service, cancel_in)

    def cancelling(*a, **k):
        service.cancel_indexing()
        return real(*a, **k)

    async def acancelling(*a, **k):
        service.cancel_indexing()
        return await real(*a, **k)

    setattr(service, cancel_in, acancelling if cancel_in == "_detect_changes" else cancelling)
    ran = []

    async def pipeline(*_a, **_k):
        ran.append(1)

    service._batch_index_pipeline = pipeline
    idx.progress.status = "idle"
    try:
        await service.index_folders([str(corpus)])
    finally:
        await mgr.close()

    assert ran == []
    assert idx.progress.status == "idle"


# --- A1-16 ---------------------------------------------------------------


def test_a1_16_chunk_overlap_is_capped_and_defaults_untouched(monkeypatch):
    from app.config import settings

    assert settings.chunk_size == 1024 and settings.chunk_overlap == 102
    assert idx.IndexingService(MagicMock(), _Emb(), _Lance()).chunk_overlap == 102

    monkeypatch.setattr(settings, "chunk_overlap", 900)
    svc = idx.IndexingService(MagicMock(), _Emb(), _Lance())
    assert svc.chunk_overlap == 512
    text = ("The quick brown fox jumps over the lazy dog. " * 5 + "\n") * 50
    chunker = idx.StreamChunker(svc.chunk_size, svc.chunk_overlap, "")
    assert len(chunker.process(text) + chunker.finalize()) < 40  # was 4739


# --- A1-17 ---------------------------------------------------------------


@pytest.mark.asyncio
async def test_a1_17_vanished_file_leaves_no_ghost_error_row(tmp_path):
    import asyncio

    mgr, service = await _service(tmp_path)
    queue: asyncio.Queue = asyncio.Queue()
    try:
        await service._stream_extract_and_prepare(
            tmp_path / "gone.md", "docs", str(tmp_path), None, queue
        )
    finally:
        await mgr.close()
    assert queue.empty()


# --- A1-19 ---------------------------------------------------------------


def test_a1_19_python_scanner_skips_hidden_entries_like_rust(tmp_path):
    from app.scanner.scanner import _scandir_walk

    (tmp_path / ".git").mkdir()
    (tmp_path / ".git" / "notes.txt").write_text("x")
    (tmp_path / ".venv" / "lib").mkdir(parents=True)
    (tmp_path / ".venv" / "lib" / "site.py").write_text("x")
    (tmp_path / ".hidden_note.md").write_text("x")
    (tmp_path / "app.py").write_text("x")
    (tmp_path / "sub").mkdir()
    (tmp_path / "sub" / "b.md").write_text("x")

    found = {
        p.relative_to(tmp_path).as_posix() for p in _scandir_walk(tmp_path, {".py", ".md", ".txt"})
    }
    assert found == {"app.py", "sub/b.md"}


# --- A1-20 ---------------------------------------------------------------


@pytest.mark.asyncio
async def test_a1_20_failed_run_keeps_failed_label_and_real_cause(tmp_path):
    corpus = tmp_path / "docs"
    corpus.mkdir()
    (corpus / "a.txt").write_text("hello world " * 20, encoding="utf-8")
    mgr, service = await _service(tmp_path)

    async def crashing_pipeline(*_a, **_k):
        raise ExceptionGroup("unhandled errors in a TaskGroup", [RuntimeError("lance add failed")])

    service._batch_index_pipeline = crashing_pipeline
    idx.progress.status = "idle"
    try:
        await service.index_folders([str(corpus)])
    finally:
        await mgr.close()

    assert idx.progress.run_failed is True
    assert idx.progress.current_file == "Failed"  # was "Checkpointing the write-ahead log…"
    assert idx.progress.last_error == "RuntimeError: lance add failed"


@pytest.mark.asyncio
async def test_a1_20_empty_run_does_not_leave_scanning_label(tmp_path):
    empty = tmp_path / "empty"
    empty.mkdir()
    mgr, service = await _service(tmp_path)
    idx.progress.status = "idle"
    try:
        await service.index_folders([str(empty)])
    finally:
        await mgr.close()
    assert idx.progress.current_file != "Scanning folders…"
    assert idx.progress.status == "idle"


# --- A1-21 ---------------------------------------------------------------


def test_a1_21_epub_entities_are_decoded_and_inline_tags_keep_words_whole(tmp_path):
    import zipfile

    from app.indexing.extractors.epub_extractor import EpubExtractor

    body = (
        "<html><body><p>Don&#8217;t trust AT&amp;T&#x2019;s caf&eacute; menu, "
        "<b>W</b>ord and <em>emph</em>asis really matter here for length.</p>"
        "<p>Second paragraph.</p></body></html>"
    )
    path = tmp_path / "b.epub"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("c.xhtml", body)
    out = " ".join(EpubExtractor().extract_stream(path, 100_000))
    assert "Don\u2019t" in out
    assert "AT&T\u2019s caf\u00e9" in out
    assert "Word" in out and "emphasis" in out
    assert "Second paragraph." in out


# --- A1-13 x A1-01 interaction -------------------------------------------


@pytest.mark.asyncio
async def test_a1_13_json_windows_get_no_injected_separator(tmp_path):
    import json

    corpus = tmp_path / "docs"
    corpus.mkdir()
    # One 300k-char token: pretty-printed it spans 3 x 128k-char windows, so a
    # separator added between windows lands inside the token as "a\na".
    path = corpus / "blob.json"
    path.write_text(json.dumps(["a" * 300_000]), encoding="utf-8")
    assert path.stat().st_size < 500_000
    (corpus / "people.csv").write_text("name,city\nAlice,Paris\nBob,Oslo\n", encoding="utf-8")

    mgr, service = await _service(tmp_path)
    idx.progress.status = "idle"
    try:
        await service.index_folders([str(corpus)])
        text = "\n".join(await _chunk_texts(mgr))
    finally:
        await mgr.close()

    assert "a\na" not in text
    assert "Paris\nname: Bob" in text  # record-style extractors still separated


# --- A2-01 ---------------------------------------------------------------


@pytest.mark.asyncio
async def test_a2_01_failed_lancedb_add_does_not_leave_files_marked_indexed(tmp_path):
    from app.storage.db import DatabaseManager
    from app.vector_store.lancedb_client import LanceDBClient

    corpus = tmp_path / "docs"
    corpus.mkdir()
    for n in ("a", "b", "c"):
        (corpus / f"{n}.txt").write_text((f"{n}word " * 80), encoding="utf-8")

    mgr = DatabaseManager(str(tmp_path / "a201.db"))
    await mgr.init_db(schema_path="app/storage/schema.sql")
    lance = LanceDBClient(persist_directory=str(tmp_path / "lance"))
    real_add = lance.add_documents
    calls = {"n": 0}

    async def flaky_add(ids, embs, metas):
        calls["n"] += 1
        if calls["n"] == 1:
            raise OSError("disk full")
        await real_add(ids, embs, metas)

    lance.add_documents = flaky_add  # type: ignore[method-assign]
    service = idx.IndexingService(mgr, _Emb(), lance)  # type: ignore[arg-type]
    try:
        idx.progress.status = "idle"
        await service.index_folders([str(corpus)])  # run 1: the add fails once
        idx.progress.status = "idle"
        await service.index_folders([str(corpus)])  # run 2: must recover
        chunk_ids = {str(r[0]) for r in await mgr.execute_query("SELECT id FROM chunks")}
        assert chunk_ids
        assert lance.get_all_ids() == chunk_ids
    finally:
        await mgr.close()


# --- A2-04 ---------------------------------------------------------------


@pytest.mark.asyncio
async def test_a2_04_reembed_refuses_while_indexing_and_holds_the_lock_itself():
    from app.indexing.reembed import ReembedError, reembed_all

    cleared = []
    during = []

    class _L:
        async def clear_all(self):
            cleared.append(1)
            during.append(idx.indexing_lock.locked())
            raise RuntimeError("stop after the first write")

    async with idx.indexing_lock:
        with pytest.raises(ReembedError):
            await reembed_all(MagicMock(), None, _L())
    assert cleared == []  # nothing was dropped while an index run held the lock

    with pytest.raises(RuntimeError, match="stop after"):
        await reembed_all(MagicMock(), None, _L())
    assert during == [True]  # the re-embed itself excludes index runs and the watcher


# --- A6-13 ---------------------------------------------------------------


@pytest.mark.asyncio
async def test_a6_13_pdf_indexed_before_ocr_was_enabled_is_gated_later(tmp_path, monkeypatch):
    from app.config import settings
    from app.indexing.extractors import ExtractMeta
    from app.indexing.extractors.pdf_extractor import PdfExtractor

    state = {"meta": ExtractMeta(page_count=3, ocr_pages=()), "calls": 0}

    def fake_stream(self, path, max_file_size):
        state["calls"] += 1
        yield "native cover page text " * 40
        if state["meta"] is not None:
            yield state["meta"]

    monkeypatch.setattr(PdfExtractor, "extract_stream", fake_stream)
    monkeypatch.setattr(settings, "ocr_enabled", True)
    monkeypatch.setattr(settings, "ocr_tier", "cpu")

    corpus = tmp_path / "docs"
    corpus.mkdir()
    (corpus / "mixed.pdf").write_bytes(b"%PDF-1.4 stand-in, the extractor is patched")

    mgr, service = await _service(tmp_path)
    try:
        idx.progress.status = "idle"
        await service.index_folders([str(corpus)])
        q = "SELECT extract_status FROM files"
        assert [r[0] for r in await mgr.execute_query(q)] == ["ocr_gated"]

        # The row as an older build stored it: indexed with OCR off, so no gate ran.
        await mgr.execute_write("UPDATE files SET extract_status = ''")
        state["meta"] = ExtractMeta(page_count=3, ocr_pages=(1, 2))
        before = state["calls"]
        idx.progress.status = "idle"
        await service.index_folders([str(corpus)])  # file unchanged -> skipped by change detection

        assert state["calls"] == before + 1
        queued = await mgr.execute_query("SELECT pages_json FROM ocr_queue")
        assert [r[0] for r in queued] == ["[1, 2]"]
        assert [r[0] for r in await mgr.execute_query(q)] == ["ocr_gated"]

        idx.progress.status = "idle"
        await service.index_folders([str(corpus)])
        assert state["calls"] == before + 1  # gated once, not on every run
    finally:
        await mgr.close()


# --- A9-12 ---------------------------------------------------------------


@pytest.mark.asyncio
async def test_a9_12_lancedb_flush_does_not_run_a_full_gc_on_the_loop(monkeypatch):
    import gc

    calls = []
    monkeypatch.setattr(gc, "collect", lambda *a, **k: calls.append(1) or 0)
    service = idx.IndexingService(MagicMock(), _Emb(), _Lance())
    await service._flush_pending_chunks_lancedb(["1"], [np.zeros(2)], [{"chunk_id": "1"}])
    assert calls == []
